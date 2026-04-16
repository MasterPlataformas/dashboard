try:
    import sys
    if sys.stdout:
        sys.stdout.reconfigure(encoding='utf-8')
except:
    pass

import importlib.util
import os
import sys
import json
import ctypes

# ============================================================
# CAMINHO DINAMICO
# ============================================================

def caminho_base():
    if getattr(sys, 'frozen', False):
        return os.path.dirname(sys.executable)
    return os.path.dirname(os.path.abspath(__file__))

PASTA_BASE     = caminho_base()
DRIVER_PATH    = os.path.join(PASTA_BASE, "msedgedriver.exe")
ARQUIVO_SAIDA  = os.path.join(PASTA_BASE, "vagas_porto_seguro.xlsx")
ARQUIVO_CONFIG = os.path.join(PASTA_BASE, "config.json")

# ============================================================
# CONFIGURACOES PADRAO
# ============================================================

CONFIG_PADRAO = {
    "usuario": "214.972.655-68",
    "senha": "Cortina10*",
    "susep": "U40TPJ (P)",
    "email_remetente": "erik.carvalhoadm@gmail.com",
    "email_senha_app": "COLE_SUA_SENHA_DE_APP_AQUI",
    "email_destinatario": "erik.carvalhoadm@gmail.com",
    "grupos_monitorados": "",
}

URL_LOGIN     = "https://corretor.portoseguro.com.br/portal/site/corretoronline/template.LOGIN/"
URL_SIMULADOR = "https://areaparceiro.portoseguro.com.br/simulador-vendas"
TIMEOUT = 40

TIPOS_DE_BEM = ["Todos os tipos", "Automoveis", "Veiculos pesados", "Imoveis"]

FAIXAS_POR_TIPO = {
    "Automoveis": [
        "De R$ 25.000,00 ate R$ 50.000,00",
        "De R$ 34.000,00 ate R$ 65.000,00",
        "De R$ 62.500,00 ate R$ 125.000,00",
        "De R$ 125.000,00 ate R$ 200.000,00",
    ],
    "Veiculos pesados": ["De R$ 180.000,00 ate R$ 360.000,00"],
    "Imoveis": [
        "De R$ 70.000,00 ate R$ 140.000,00",
        "De R$ 140.000,00 ate R$ 280.000,00",
        "De R$ 280.000,00 ate R$ 560.000,00",
        "De R$ 600.000,00 ate R$ 900.000,00",
        "De R$ 900.000,00 ate R$ 1.000.000,00",
    ],
}
FAIXAS_POR_TIPO["Todos os tipos"] = (
    [f"[Auto] {f}" for f in FAIXAS_POR_TIPO["Automoveis"]] +
    [f"[Pesados] {f}" for f in FAIXAS_POR_TIPO["Veiculos pesados"]] +
    [f"[Imoveis] {f}" for f in FAIXAS_POR_TIPO["Imoveis"]]
)

FAIXAS_REAIS = {
    "Automoveis": [
        "De R$ 25.000,00 até R$ 50.000,00",
        "De R$ 34.000,00 até R$ 65.000,00",
        "De R$ 62.500,00 até R$ 125.000,00",
        "De R$ 125.000,00 até R$ 200.000,00",
    ],
    "Veiculos pesados": ["De R$ 180.000,00 até R$ 360.000,00"],
    "Imoveis": [
        "De R$ 70.000,00 até R$ 140.000,00",
        "De R$ 140.000,00 até R$ 280.000,00",
        "De R$ 280.000,00 até R$ 560.000,00",
        "De R$ 600.000,00 até R$ 900.000,00",
        "De R$ 900.000,00 até R$ 1.000.000,00",
    ],
}

TIPOS_REAIS = {
    "Automoveis": "Automóveis",
    "Veiculos pesados": "Veículos pesados",
    "Imoveis": "Imóveis",
}

# ============================================================
# PALETA DE CORES
# ============================================================

COR = {
    "primary":      "#1F3864",
    "primary_dark": "#152848",
    "primary_light":"#2A4F8F",
    "accent":       "#0046C0",
    "success":      "#1B5E20",
    "success_light":"#E8F5E9",
    "danger":       "#B71C1C",
    "danger_light": "#FFEBEE",
    "warning_bg":   "#FFF8E1",
    "warning_brd":  "#FFE082",
    "monitor_bg":   "#EEF3FC",
    "monitor_brd":  "#C5D5F0",
    "bg":           "#F5F7FA",
    "surface":      "#FFFFFF",
    "surface2":     "#F0F4F8",
    "border":       "#D0D8E8",
    "text":         "#1A1A2E",
    "text_muted":   "#5A6A85",
    "log_bg":       "#0F1117",
    "log_fg":       "#CDD6F4",
    "tab_bg":       "#E8EDF5",
}

# ============================================================
# CONFIG
# ============================================================

def carregar_config():
    if os.path.exists(ARQUIVO_CONFIG):
        try:
            with open(ARQUIVO_CONFIG, "r", encoding="utf-8") as f:
                dados = json.load(f)
                config = CONFIG_PADRAO.copy()
                config.update(dados)
                return config
        except:
            pass
    return CONFIG_PADRAO.copy()

def salvar_config(config):
    with open(ARQUIVO_CONFIG, "w", encoding="utf-8") as f:
        json.dump(config, f, ensure_ascii=False, indent=2)


def deduplicar_grupos_encontrados(grupos_encontrados):
    vistos = set()
    unicos = []
    for grupo in grupos_encontrados:
        chave = grupo.get("grupo", "").strip().upper()
        if not chave or chave in vistos:
            continue
        vistos.add(chave)
        unicos.append(grupo)
    return unicos

# ============================================================
# IMPORTS
# ============================================================

import threading
import tkinter as tk
from tkinter import ttk, messagebox
from selenium import webdriver
from selenium.webdriver.edge.service import Service as EdgeService
from selenium.webdriver.edge.webdriver import WebDriver as EdgeDriver
from selenium.webdriver.edge.options import Options as EdgeOptions
from selenium.webdriver.common.by import By
from selenium.webdriver.support.ui import WebDriverWait, Select
from selenium.webdriver.support import expected_conditions as EC
from selenium.webdriver.common.keys import Keys
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment
from datetime import datetime
import time
import smtplib
from email.mime.text import MIMEText
from email.mime.multipart import MIMEMultipart

# ============================================================
# EMAIL
# ============================================================

def enviar_email(grupos_encontrados, config):
    try:
        grupos_encontrados = deduplicar_grupos_encontrados(grupos_encontrados)
        msg = MIMEMultipart()
        msg["From"] = config["email_remetente"]
        msg["To"] = config["email_destinatario"]
        msg["Subject"] = "Alerta Porto Seguro - Grupos Disponiveis!"
        linhas_html = "".join(
            f"<tr><td style='padding:8px;border:1px solid #ddd'><b>{g['grupo']}</b></td>"
            f"<td style='padding:8px;border:1px solid #ddd'>{g['faixa']}</td></tr>"
            for g in grupos_encontrados
        )
        corpo = f"""<html><body>
        <h2 style='color:#1F3864'>Grupos Monitorados Encontrados!</h2>
        <table style='border-collapse:collapse;width:100%'>
            <tr style='background:#1F3864;color:white'>
                <th style='padding:8px'>Grupo</th><th style='padding:8px'>Faixa</th>
            </tr>{linhas_html}
        </table>
        <p style='color:#888;font-size:12px'>Verificado em {datetime.now().strftime('%d/%m/%Y as %H:%M:%S')}</p>
        </body></html>"""
        msg.attach(MIMEText(corpo, "html"))
        with smtplib.SMTP_SSL("smtp.gmail.com", 465) as smtp:
            smtp.login(config["email_remetente"], config["email_senha_app"])
            smtp.send_message(msg)
        return True
    except Exception as e:
        return str(e)

# ============================================================
# SELENIUM
# ============================================================

def configurar_navegador(exibir_janela=True):
    options = EdgeOptions()
    options.add_argument("--disable-gpu")
    options.add_argument("--no-sandbox")
    options.add_argument("--disable-dev-shm-usage")
    options.add_argument("--disable-blink-features=AutomationControlled")
    if not exibir_janela:
        options.add_argument("--headless=new")
    driver = webdriver.Edge(options=options)
    driver.set_page_load_timeout(60)
    driver.set_window_size(580, 780)
    if exibir_janela:
        pos_x, pos_y = obter_posicao_janela_extrator()
        driver.set_window_position(pos_x, pos_y)
    return driver

def obter_posicao_janela_extrator():
    try:
        user32 = ctypes.windll.user32
        largura = user32.GetSystemMetrics(0)
        pos_x = max(24, min(40, largura - 620))
        return pos_x, 24
    except Exception:
        return 40, 24


def fazer_login(driver, wait, log, config):
    espera = 5
    log("Acessando pagina de login...")
    driver.get(URL_LOGIN)
    time.sleep(espera)
    try:
        botao = wait.until(EC.element_to_be_clickable(
            (By.XPATH, "//*[contains(text(), 'ACESSAR O CORRETOR ONLINE')]")
        ))
        driver.execute_script("arguments[0].click();", botao)
        time.sleep(espera)
        log("Clicou em Acessar o Corretor Online")
    except:
        pass
    time.sleep(3)
    campo_user = wait.until(EC.element_to_be_clickable((By.ID, "logonPrincipal")))
    driver.execute_script("arguments[0].scrollIntoView(true);", campo_user)
    campo_user.click()
    time.sleep(0.3)
    for c in config["usuario"]:
        campo_user.send_keys(c)
        time.sleep(0.05)
    campo_pass = driver.find_element(By.NAME, "password")
    campo_pass.click()
    time.sleep(0.3)
    for c in config["senha"]:
        campo_pass.send_keys(c)
        time.sleep(0.05)
    time.sleep(0.5)
    driver.find_element(By.ID, "inputLogin").click()
    log("Login enviado...")
    time.sleep(espera + 3)
    try:
        dropdown = wait.until(EC.presence_of_element_located(
            (By.XPATH, "//select[contains(@class,'susep') or @id='susep' or @name='susep']")
        ))
        Select(dropdown).select_by_visible_text(config["susep"])
        time.sleep(1)
        driver.find_element(By.ID, "btnAvancarSusep").click()
        time.sleep(espera)
        log("Login completo!")
    except Exception as e:
        try:
            screenshot_path = os.path.join(PASTA_BASE, "debug_susep.png")
            driver.save_screenshot(screenshot_path)
            log(f"Screenshot salvo em: {screenshot_path}")
        except:
            pass
        log(f"SUSEP nao encontrada: {str(e)[:80]}")


def acessar_simulador_faixa(driver, wait, log, tipo_chave, faixa_credito):
    tipo_real = TIPOS_REAIS.get(tipo_chave, tipo_chave)
    log(f"--- {tipo_real} | {faixa_credito} ---")
    driver.get(URL_SIMULADOR)
    time.sleep(6)
    try:
        scrim = driver.find_element(By.CSS_SELECTOR, ".v-navigation-drawer__scrim")
        driver.execute_script("arguments[0].click();", scrim)
        time.sleep(1)
    except:
        pass
    card = wait.until(EC.presence_of_element_located(
        (By.XPATH, f"//div[contains(@class,'card-text') and contains(text(),'{tipo_real}')]")
    ))
    driver.execute_script("arguments[0].click();", card)
    time.sleep(3)
    dropdown = wait.until(EC.presence_of_element_located(
        (By.XPATH, "//div[contains(@class,'porto-select-field')]")
    ))
    driver.execute_script("arguments[0].click();", dropdown)
    time.sleep(1)
    opcao = wait.until(EC.presence_of_element_located(
        (By.XPATH, f"//*[contains(text(),'{faixa_credito}')]")
    ))
    driver.execute_script("arguments[0].click();", opcao)
    time.sleep(2)
    for texto_grupo in ["Tipo de grupo", "Tipo de cliente"]:
        drop = wait.until(EC.presence_of_element_located(
            (By.XPATH, f"//div[contains(@class,'porto-select-field') and .//*[contains(text(),'{texto_grupo}')]]")
        ))
        driver.execute_script("arguments[0].click();", drop)
        time.sleep(1)
        opcao_texto = "Em andamento" if texto_grupo == "Tipo de grupo" else "Pessoa física"
        op = wait.until(EC.presence_of_element_located(
            (By.XPATH, f"//*[contains(text(),'{opcao_texto}')]")
        ))
        driver.execute_script("arguments[0].click();", op)
        time.sleep(1)
    driver.execute_script("document.body.click();")
    time.sleep(1)
    botao_ver = driver.find_element(By.CSS_SELECTOR, "button.default-button.primary")
    driver.execute_script("arguments[0].scrollIntoView(true);", botao_ver)
    driver.execute_script("arguments[0].click();", botao_ver)
    time.sleep(5)
    try:
        driver.execute_script("window.scrollTo(0, document.body.scrollHeight);")
        time.sleep(1)
        icone = wait.until(EC.presence_of_element_located((
            By.XPATH, "//i[contains(@class,'v-select__menu-icon')]"
        )))
        driver.execute_script("arguments[0].scrollIntoView({block:'center'});", icone)
        time.sleep(1)
        driver.execute_script("""
            var el = arguments[0];
            ['mousedown','mouseup','click'].forEach(function(evt) {
                el.dispatchEvent(new MouseEvent(evt, {bubbles:true, cancelable:true, view:window}));
            });
        """, icone)
        time.sleep(2)
        opcoes = wait.until(EC.presence_of_all_elements_located((
            By.XPATH, "//div[@role='listbox']//div[@role='option']"
        )))
        opcao_50 = next((op for op in opcoes if op.text.strip() == "50"), None)
        if opcao_50:
            driver.execute_script("""
                var el = arguments[0];
                ['mousedown','mouseup','click'].forEach(function(evt) {
                    el.dispatchEvent(new MouseEvent(evt, {bubbles:true, cancelable:true, view:window}));
                });
            """, opcao_50)
            time.sleep(2)
            wait.until(lambda d: len(d.find_elements(By.XPATH, "//table//tr")) > 11)
            log("  Exibindo 50 itens por pagina.")
    except Exception as e:
        log(f"  Aviso paginacao: {str(e)[:60]}")
    linhas = []
    cabecalho = []
    pagina = 1
    while True:
        log(f"  Pagina {pagina}...")
        wait.until(EC.presence_of_element_located((By.TAG_NAME, "table")))
        time.sleep(1)
        tabela = driver.find_element(By.TAG_NAME, "table")
        if pagina == 1:
            ths = tabela.find_elements(By.TAG_NAME, "th")
            cabecalho = [th.text.strip() for th in ths if th.text.strip()]
        trs = tabela.find_elements(By.TAG_NAME, "tr")
        count = 0
        for tr in trs:
            tds = tr.find_elements(By.TAG_NAME, "td")
            texto = [td.text.strip() for td in tds]
            if any(texto):
                linhas.append(texto)
                count += 1
        log(f"    {count} linhas")
        try:
            botoes = driver.find_elements(By.XPATH,
                "//button[contains(@class,'default-button secondary flat ml-6') and not(@disabled)]")
            if not botoes:
                break
            proxima = botoes[-1]
            driver.execute_script("arguments[0].scrollIntoView({block:'center'});", proxima)
            driver.execute_script("arguments[0].click();", proxima)
            time.sleep(2)
            pagina += 1
        except:
            break
    log(f"  Total: {len(linhas)} linhas")
    return cabecalho, linhas


def salvar_excel(abas, log, arquivo=None):
    if arquivo is None:
        arquivo = ARQUIVO_SAIDA
    log("Salvando Excel...")
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    for nome_aba, cabecalho, dados in abas:
        ws = wb.create_sheet(title=nome_aba[:31])
        todas = [cabecalho] + dados if cabecalho else dados
        for i, linha in enumerate(todas):
            for j, valor in enumerate(linha):
                celula = ws.cell(row=i+1, column=j+1, value=valor)
                celula.alignment = Alignment(horizontal="left", vertical="center")
                if i == 0:
                    celula.font = Font(bold=True, color="FFFFFF", size=11)
                    celula.fill = PatternFill("solid", fgColor="1F3864")
                elif i % 2 == 0:
                    celula.fill = PatternFill("solid", fgColor="DCE6F1")
        for col in ws.columns:
            max_len = max((len(str(c.value)) for c in col if c.value), default=10)
            ws.column_dimensions[col[0].column_letter].width = min(max_len + 4, 50)
    ws_log = wb.create_sheet(title="Informacoes")
    ws_log["A1"] = "Extracao realizada em:"
    ws_log["B1"] = datetime.now().strftime("%d/%m/%Y as %H:%M:%S")
    ws_log["A1"].font = Font(bold=True)
    wb.save(arquivo)
    log(f"Arquivo salvo: {arquivo}")

# ============================================================
# ESTILOS TTK
# ============================================================

def aplicar_estilos():
    style = ttk.Style()
    style.theme_use("clam")

    # Notebook
    style.configure("TNotebook",
        background=COR["bg"],
        borderwidth=0,
        tabmargins=[0, 4, 0, 0],
    )
    style.configure("TNotebook.Tab",
        background=COR["tab_bg"],
        foreground=COR["text_muted"],
        font=("Segoe UI", 9, "bold"),
        padding=[16, 8],
        borderwidth=0,
    )
    style.map("TNotebook.Tab",
        background=[("selected", COR["surface"]), ("active", COR["border"])],
        foreground=[("selected", COR["primary"])],
    )

    # Combobox
    style.configure("TCombobox",
        fieldbackground=COR["surface"],
        background=COR["surface"],
        foreground=COR["text"],
        bordercolor=COR["border"],
        arrowcolor=COR["accent"],
        padding=6,
    )
    style.map("TCombobox",
        fieldbackground=[("readonly", COR["surface"])],
        selectbackground=[("readonly", COR["surface"])],
        selectforeground=[("readonly", COR["text"])],
    )

    # Scrollbar
    style.configure("TScrollbar",
        background=COR["border"],
        troughcolor=COR["bg"],
        borderwidth=0,
        arrowsize=12,
    )

# ============================================================
# WIDGETS CUSTOMIZADOS
# ============================================================

def label_secao(parent, texto, bg=None):
    bg = bg or COR["bg"]
    f = tk.Frame(parent, bg=bg)
    f.pack(fill="x", pady=(16, 4))
    tk.Label(f, text=texto, font=("Segoe UI", 10, "bold"),
             bg=bg, fg=COR["primary"]).pack(anchor="w")
    tk.Frame(f, bg=COR["border"], height=1).pack(fill="x", pady=(4, 0))
    return f

def card(parent, bg=None, pad_x=12, pad_y=10):
    bg = bg or COR["surface"]
    f = tk.Frame(parent, bg=bg, highlightbackground=COR["border"],
                 highlightthickness=1)
    f.pack(fill="x", pady=(0, 10))
    inner = tk.Frame(f, bg=bg, padx=pad_x, pady=pad_y)
    inner.pack(fill="x")
    return inner

def botao_primario(parent, texto, comando, bg=None, pady=10):
    bg = bg or COR["accent"]
    b = tk.Button(parent, text=texto, font=("Segoe UI", 10, "bold"),
                  bg=bg, fg="white", relief="flat",
                  activebackground=COR["primary_dark"], activeforeground="white",
                  cursor="hand2", pady=pady, command=comando,
                  bd=0, highlightthickness=0)
    b.pack(fill="x", pady=(0, 6))
    return b

def botao_secundario(parent, texto, comando, bg=None, state="normal"):
    bg = bg or COR["monitor_bg"]
    b = tk.Button(parent, text=texto, font=("Segoe UI", 9),
                  bg=bg, fg=COR["primary"], relief="flat",
                  activebackground=COR["border"], cursor="hand2",
                  pady=6, command=comando, state=state,
                  bd=0, highlightthickness=0)
    b.pack(fill="x", pady=(0, 6))
    return b

# ============================================================
# INTERFACE
# ============================================================

class App:
    def __init__(self, root, mostrar_navegador_inicial=True):
        self.root = root
        self.root.title("Extrator de Vagas — Porto Seguro")
        self.root.geometry("680x860+0+0")
        self.root.minsize(640, 780)
        self.root.resizable(True, True)
        self.root.configure(bg=COR["bg"])

        aplicar_estilos()

        self.config = carregar_config()
        self.rodando = False
        self.monitorando = False
        self.faixa_vars = {}
        self.parar_monitor = threading.Event()
        self.cancelar_execucao = threading.Event()
        self.mostrar_navegador_inicial = mostrar_navegador_inicial
        self.dashboard_window = None
        self.driver_ativo = None
        self._save_grupos_job = None

        self._build_ui()
        self.root.protocol("WM_DELETE_WINDOW", self.fechar_aplicacao)

    def _build_ui(self):
        # Header
        header = tk.Frame(self.root, bg=COR["primary"], pady=16)
        header.pack(fill="x", side="top")

        tk.Label(header, text="Extrator de Vagas",
                 font=("Segoe UI", 16, "bold"),
                 bg=COR["primary"], fg="white").pack()
        tk.Label(header, text="Porto Seguro Consorcios",
                 font=("Segoe UI", 9),
                 bg=COR["primary"], fg="#A8C4E0").pack()

        # Status bar
        self.status_bar = tk.Frame(self.root, bg=COR["primary_dark"], pady=4)
        self.status_bar.pack(fill="x", side="top")
        self.label_status = tk.Label(self.status_bar, text="Pronto",
                                      font=("Segoe UI", 8),
                                      bg=COR["primary_dark"], fg="#A8C4E0")
        self.label_status.pack(side="left", padx=12)

        # Log fixo na base
        log_outer = tk.Frame(self.root, bg=COR["bg"], padx=16, pady=8)
        log_outer.pack(fill="x", side="bottom")
        tk.Label(log_outer, text="Log de execucao",
                 font=("Segoe UI", 8, "bold"),
                 bg=COR["bg"], fg=COR["text_muted"]).pack(anchor="w", pady=(0, 4))
        log_inner = tk.Frame(log_outer, bg=COR["log_bg"],
                              highlightbackground=COR["border"], highlightthickness=1)
        log_inner.pack(fill="x")
        self.log_box = tk.Text(log_inner, height=7, font=("Consolas", 9),
                                bg=COR["log_bg"], fg=COR["log_fg"], relief="flat",
                                state="disabled", wrap="word", padx=8, pady=6,
                                insertbackground=COR["log_fg"])
        scroll_log = ttk.Scrollbar(log_inner, command=self.log_box.yview)
        self.log_box.configure(yscrollcommand=scroll_log.set)
        self.log_box.pack(side="left", fill="x", expand=True)
        scroll_log.pack(side="right", fill="y")

        # Notebook
        self.notebook = ttk.Notebook(self.root)
        self.notebook.pack(fill="both", expand=True, side="top")

        self.aba_extracao = tk.Frame(self.notebook, bg=COR["bg"])
        self.aba_config   = tk.Frame(self.notebook, bg=COR["bg"])
        self.notebook.add(self.aba_extracao, text="  Extracao  ")
        self.notebook.add(self.aba_config,   text="  Configuracoes  ")

        self._build_aba_extracao()
        self._build_aba_config()

    def _build_aba_extracao(self):
        canvas = tk.Canvas(self.aba_extracao, bg=COR["bg"], highlightthickness=0)
        scrollbar = ttk.Scrollbar(self.aba_extracao, orient="vertical", command=canvas.yview)
        canvas.configure(yscrollcommand=scrollbar.set)
        scrollbar.pack(side="right", fill="y")
        canvas.pack(side="left", fill="both", expand=True)

        inner = tk.Frame(canvas, bg=COR["bg"])
        cw = canvas.create_window((0, 0), window=inner, anchor="nw")

        def on_cfg(e):
            canvas.configure(scrollregion=canvas.bbox("all"))
            canvas.itemconfig(cw, width=canvas.winfo_width())

        inner.bind("<Configure>", on_cfg)
        canvas.bind("<Configure>", lambda e: canvas.itemconfig(cw, width=e.width))
        canvas.bind_all("<MouseWheel>", lambda e: canvas.yview_scroll(-1*(e.delta//120), "units"))

        frame = tk.Frame(inner, bg=COR["bg"], padx=20, pady=16)
        frame.pack(fill="x")

        hero = card(frame, bg=COR["surface"], pad_x=16, pad_y=14)
        tk.Label(hero, text="Extracao e monitoramento em um unico fluxo",
                 font=("Segoe UI", 12, "bold"),
                 bg=COR["surface"], fg=COR["text"]).pack(anchor="w")
        tk.Label(
            hero,
            text="Configure as faixas, acompanhe o log em tempo real e abra o dashboard sem interromper a monitoracao.",
            font=("Segoe UI", 9),
            bg=COR["surface"], fg=COR["text_muted"],
            wraplength=560, justify="left"
        ).pack(anchor="w", pady=(6, 0))

        # Tipo de bem
        tipo_card = card(frame)
        label_secao(tipo_card, "Tipo de bem", bg=COR["surface"])
        self.tipo_bem = ttk.Combobox(tipo_card, values=TIPOS_DE_BEM, state="readonly",
                                      font=("Segoe UI", 10), width=50)
        self.tipo_bem.current(0)
        self.tipo_bem.pack(fill="x", pady=(0, 4), ipady=4)
        self.tipo_bem.bind("<<ComboboxSelected>>", self.atualizar_faixas)

        # Faixas
        faixas_card = card(frame)
        label_secao(faixas_card, "Faixas de credito", bg=COR["surface"])
        self.frame_faixas = tk.Frame(faixas_card, bg=COR["surface"],
                                      highlightbackground=COR["border"], highlightthickness=1)
        self.frame_faixas.pack(fill="x", pady=(0, 2))
        self.atualizar_faixas()

        # Grupos
        grupos_card = card(frame)
        label_secao(grupos_card, "Grupos para monitorar", bg=COR["surface"])
        tk.Label(grupos_card, text="Digite um grupo por linha. Um alerta sera exibido se algum for encontrado.",
                 font=("Segoe UI", 8), bg=COR["surface"], fg=COR["text_muted"],
                 wraplength=560, justify="left").pack(anchor="w", pady=(0, 6))
        self.grupos_box = tk.Text(grupos_card, height=4, font=("Consolas", 10),
                                   bg=COR["surface"], fg=COR["text"], relief="flat",
                                   highlightbackground=COR["border"], highlightthickness=1,
                                   padx=8, pady=6)
        self.grupos_box.pack(fill="x", pady=(0, 2))
        grupos_salvos = self.config.get("grupos_monitorados", "").strip()
        if grupos_salvos:
            self.grupos_box.insert("1.0", grupos_salvos)
        self.grupos_box.bind("<KeyRelease>", self._agendar_salvar_grupos_monitorados)
        self.grupos_box.bind("<FocusOut>", self._salvar_grupos_monitorados)

        # Monitoramento
        label_secao(frame, "Monitoramento automatico")
        mon_card = tk.Frame(frame, bg=COR["monitor_bg"],
                             highlightbackground=COR["monitor_brd"], highlightthickness=1)
        mon_card.pack(fill="x", pady=(0, 12))
        mon_inner = tk.Frame(mon_card, bg=COR["monitor_bg"], padx=14, pady=12)
        mon_inner.pack(fill="x")

        resumo = tk.Frame(mon_inner, bg=COR["surface"], padx=12, pady=10,
                          highlightbackground=COR["monitor_brd"], highlightthickness=1)
        resumo.pack(fill="x", pady=(0, 10))
        tk.Label(resumo, text="Fluxo recomendado",
                 font=("Segoe UI", 9, "bold"),
                 bg=COR["surface"], fg=COR["primary"]).pack(anchor="w")
        tk.Label(
            resumo,
            text="O botao combinado abre o dashboard e mantem a extracao ativa abaixo, dentro desta ferramenta de monitoramento.",
            font=("Segoe UI", 8),
            bg=COR["surface"], fg=COR["text_muted"],
            wraplength=540, justify="left"
        ).pack(anchor="w", pady=(4, 0))

        # Intervalo
        row_int = tk.Frame(mon_inner, bg=COR["monitor_bg"])
        row_int.pack(fill="x", pady=(0, 8))
        tk.Label(row_int, text="Verificar a cada:",
                 font=("Segoe UI", 9), bg=COR["monitor_bg"], fg=COR["text"]).pack(side="left")
        self.intervalo_var = tk.StringVar(value="30")
        tk.Entry(row_int, textvariable=self.intervalo_var,
                 width=5, font=("Segoe UI", 10), justify="center",
                 bg=COR["surface"], relief="flat",
                 highlightbackground=COR["border"], highlightthickness=1).pack(side="left", padx=8, ipady=3)
        self.unidade_var = tk.StringVar(value="minutos")
        ttk.Combobox(row_int, textvariable=self.unidade_var,
                     values=["minutos", "horas"], state="readonly",
                     width=9, font=("Segoe UI", 9)).pack(side="left")

        # Comportamento
        tk.Label(mon_inner, text="Ao encontrar grupo monitorado:",
                 font=("Segoe UI", 9, "bold"), bg=COR["monitor_bg"], fg=COR["text"]).pack(anchor="w", pady=(0, 4))
        self.comportamento_var = tk.StringVar(value="continuar")
        row_radio = tk.Frame(mon_inner, bg=COR["monitor_bg"])
        row_radio.pack(fill="x", pady=(0, 10))
        for txt, val in [("Continuar monitorando", "continuar"), ("Parar apos encontrar", "parar")]:
            tk.Radiobutton(row_radio, text=txt, variable=self.comportamento_var, value=val,
                           font=("Segoe UI", 9), bg=COR["monitor_bg"],
                           fg=COR["text"], activebackground=COR["monitor_bg"],
                           selectcolor=COR["surface"]).pack(side="left", padx=(0, 16))

        # Botoes monitoramento
        row_mon = tk.Frame(mon_inner, bg=COR["monitor_bg"])
        row_mon.pack(fill="x")
        self.btn_iniciar_monitor = tk.Button(row_mon, text="Iniciar monitoramento",
                                              font=("Segoe UI", 9, "bold"),
                                              bg=COR["accent"], fg="white", relief="flat",
                                              cursor="hand2", padx=14, pady=6,
                                              bd=0, command=self.iniciar_monitoramento)
        self.btn_iniciar_monitor.pack(side="left", padx=(0, 8))
        self.btn_parar_monitor = tk.Button(row_mon, text="Parar",
                                            font=("Segoe UI", 9, "bold"),
                                            bg=COR["danger"], fg="white", relief="flat",
                                            cursor="hand2", padx=14, pady=6,
                                            bd=0, command=self.parar_monitoramento,
                                            state="disabled")
        self.btn_parar_monitor.pack(side="left")
        self.btn_monitor_dashboard = tk.Button(
            row_mon, text="Monitorar + Dashboard",
            font=("Segoe UI", 9, "bold"),
            bg="#0F766E", fg="white", relief="flat",
            activebackground="#0D5E57", activeforeground="white",
            cursor="hand2", padx=14, pady=6, bd=0,
            command=self.iniciar_monitoramento_com_dashboard,
        )
        self.btn_monitor_dashboard.pack(side="left", padx=(8, 0))
        self.label_proximo = tk.Label(mon_inner, text="",
                                       font=("Segoe UI", 8), bg=COR["monitor_bg"],
                                       fg=COR["text_muted"])
        self.label_proximo.pack(anchor="w", pady=(8, 0))
        self.mostrar_navegador_var = tk.BooleanVar(value=self.mostrar_navegador_inicial)
        tk.Checkbutton(
            mon_inner,
            text="Exibir janela do navegador durante a extracao",
            variable=self.mostrar_navegador_var,
            font=("Segoe UI", 9),
            bg=COR["monitor_bg"],
            fg=COR["text"],
            activebackground=COR["monitor_bg"],
            activeforeground=COR["text"],
            selectcolor=COR["surface"],
        ).pack(anchor="w", pady=(10, 0))

        # Botoes de acao
        acoes_card = card(frame)
        label_secao(acoes_card, "Acoes", bg=COR["surface"])
        self.btn = tk.Button(acoes_card, text="EXTRAIR UMA VEZ",
                             font=("Segoe UI", 11, "bold"),
                             bg=COR["success"], fg="white", relief="flat",
                             activebackground="#145214", activeforeground="white",
                             cursor="hand2", pady=12, bd=0,
                             command=self.iniciar)
        self.btn.pack(fill="x", pady=(0, 6))

        self.btn_tudo = tk.Button(acoes_card, text="EXTRAIR TUDO  (Auto + Pesados + Imoveis)",
                                   font=("Segoe UI", 10, "bold"),
                                   bg=COR["primary"], fg="white", relief="flat",
                                   activebackground=COR["primary_dark"],
                                   cursor="hand2", pady=10, bd=0,
                                   command=self.iniciar_tudo)
        self.btn_tudo.pack(fill="x", pady=(0, 6))

        self.btn_excel = tk.Button(acoes_card, text="Abrir pasta com planilhas",
                                    font=("Segoe UI", 9),
                                    bg=COR["monitor_bg"], fg=COR["primary"], relief="flat",
                                    activebackground=COR["border"],
                                    cursor="hand2", pady=8, bd=0,
                                    command=self.abrir_pasta, state="disabled")
        self.btn_excel.pack(fill="x", pady=(0, 2))

    def _build_aba_config(self):
        canvas = tk.Canvas(self.aba_config, bg=COR["bg"], highlightthickness=0)
        scrollbar = ttk.Scrollbar(self.aba_config, orient="vertical", command=canvas.yview)
        canvas.configure(yscrollcommand=scrollbar.set)
        scrollbar.pack(side="right", fill="y")
        canvas.pack(side="left", fill="both", expand=True)
        inner = tk.Frame(canvas, bg=COR["bg"])
        cw = canvas.create_window((0, 0), window=inner, anchor="nw")

        def on_cfg(e):
            canvas.configure(scrollregion=canvas.bbox("all"))
            canvas.itemconfig(cw, width=canvas.winfo_width())

        inner.bind("<Configure>", on_cfg)
        canvas.bind("<Configure>", lambda e: canvas.itemconfig(cw, width=e.width))

        frame = tk.Frame(inner, bg=COR["bg"], padx=24, pady=20)
        frame.pack(fill="x")

        def campo(label, chave, senha=False):
            tk.Label(frame, text=label, font=("Segoe UI", 9, "bold"),
                     bg=COR["bg"], fg=COR["text"]).pack(anchor="w", pady=(12, 3))
            var = tk.StringVar(value=self.config.get(chave, ""))
            e = tk.Entry(frame, textvariable=var, font=("Segoe UI", 10),
                         relief="flat", bg=COR["surface"], fg=COR["text"],
                         highlightbackground=COR["border"], highlightthickness=1,
                         show="*" if senha else "")
            e.pack(fill="x", ipady=6)
            return var

        label_secao(frame, "Acesso ao Corretor Online")
        self.cfg_usuario = campo("CPF / Usuario:", "usuario")
        self.cfg_senha   = campo("Senha:", "senha", senha=True)
        self.cfg_susep   = campo("SUSEP:", "susep")

        label_secao(frame, "Configuracoes de E-mail")
        self.cfg_email_rem  = campo("E-mail remetente (Gmail):", "email_remetente")
        self.cfg_email_app  = campo("Senha de app do Gmail:", "email_senha_app", senha=True)
        self.cfg_email_dest = campo("E-mail destinatario:", "email_destinatario")

        tk.Label(frame, text="Como gerar a senha de app: myaccount.google.com > Seguranca > Senhas de app",
                 font=("Segoe UI", 8), bg=COR["bg"], fg=COR["text_muted"],
                 wraplength=500, justify="left").pack(anchor="w", pady=(6, 0))

        tk.Button(frame, text="Salvar configuracoes",
                  font=("Segoe UI", 11, "bold"),
                  bg=COR["accent"], fg="white", relief="flat",
                  activebackground=COR["primary_dark"],
                  cursor="hand2", pady=12, bd=0,
                  command=self.salvar_configuracoes).pack(fill="x", pady=(24, 6))

        self.label_status_config = tk.Label(frame, text="",
                                             font=("Segoe UI", 9, "bold"),
                                             bg=COR["bg"], fg=COR["success"])
        self.label_status_config.pack(anchor="w")

    def salvar_configuracoes(self):
        self.config["usuario"]            = self.cfg_usuario.get().strip()
        self.config["senha"]              = self.cfg_senha.get()
        self.config["susep"]              = self.cfg_susep.get().strip()
        self.config["email_remetente"]    = self.cfg_email_rem.get().strip()
        self.config["email_senha_app"]    = self.cfg_email_app.get().strip()
        self.config["email_destinatario"] = self.cfg_email_dest.get().strip()
        salvar_config(self.config)
        self.label_status_config.configure(text="Configuracoes salvas com sucesso!")
        self.root.after(3000, lambda: self.label_status_config.configure(text=""))

    def set_status(self, msg):
        self.label_status.configure(text=msg)
        self.root.update()

    def atualizar_faixas(self, event=None):
        for w in self.frame_faixas.winfo_children():
            w.destroy()
        self.faixa_vars.clear()
        tipo = self.tipo_bem.get()
        if tipo == "Todos os tipos":
            tk.Label(self.frame_faixas,
                     text="Todas as faixas de Automoveis, Veiculos Pesados e Imoveis serao extraidas.",
                     font=("Segoe UI", 9), bg=COR["surface"], fg=COR["text_muted"],
                     wraplength=480, justify="left").pack(anchor="w", padx=12, pady=10)
            return
        for faixa in FAIXAS_POR_TIPO.get(tipo, []):
            var = tk.BooleanVar(value=True)
            tk.Checkbutton(self.frame_faixas, text=faixa, variable=var,
                           font=("Segoe UI", 9), bg=COR["surface"], fg=COR["text"],
                           activebackground=COR["surface"], selectcolor=COR["surface"],
                           anchor="w").pack(fill="x", padx=12, pady=3)
            self.faixa_vars[faixa] = var

    def get_faixas_selecionadas(self):
        tipo = self.tipo_bem.get()
        if tipo == "Todos os tipos":
            return ["__TODOS__"]
        faixas_exibidas = FAIXAS_POR_TIPO.get(tipo, [])
        faixas_reais    = FAIXAS_REAIS.get(tipo, [])
        return [
            faixas_reais[i]
            for i, f in enumerate(faixas_exibidas)
            if f in self.faixa_vars and self.faixa_vars[f].get() and i < len(faixas_reais)
        ]

    def get_grupos_monitorados(self):
        texto = self.grupos_box.get("1.0", "end").strip()
        return [g.strip().upper() for g in texto.splitlines() if g.strip()] if texto else []

    def get_intervalo_segundos(self):
        try:
            valor = int(self.intervalo_var.get())
            return valor * 3600 if self.unidade_var.get() == "horas" else valor * 60
        except:
            return 1800

    def log(self, msg):
        self.log_box.configure(state="normal")
        self.log_box.insert("end", f"  {msg}\n")
        self.log_box.see("end")
        self.log_box.configure(state="disabled")
        self.root.update()

    def _registrar_driver(self, driver):
        self.driver_ativo = driver

    def _encerrar_driver_ativo(self):
        driver = self.driver_ativo
        self.driver_ativo = None
        if driver:
            try:
                driver.quit()
            except:
                pass

    def _agendar_salvar_grupos_monitorados(self, event=None):
        if self._save_grupos_job:
            try:
                self.root.after_cancel(self._save_grupos_job)
            except Exception:
                pass
        self._save_grupos_job = self.root.after(300, self._salvar_grupos_monitorados)

    def _salvar_grupos_monitorados(self, event=None):
        self._save_grupos_job = None
        if not hasattr(self, "grupos_box"):
            return
        self.config["grupos_monitorados"] = self.grupos_box.get("1.0", "end").strip()
        salvar_config(self.config)

    def _erro_fatal_sem_retry(self, err):
        msg = str(err).lower()
        termos_fatais = [
            "invalid session id",
            "no such window",
            "target window already closed",
            "web view not found",
            "chrome not reachable",
            "disconnected",
            "session deleted because of page crash",
            "session deleted as the browser has closed",
        ]
        return any(termo in msg for termo in termos_fatais)

    def cancelar_tudo(self, atualizar_status=True):
        self.cancelar_execucao.set()
        self.parar_monitor.set()
        self.monitorando = False
        if atualizar_status:
            self.label_proximo.configure(text="Parada solicitada...")
            self.set_status("Parando...")
            self.log("Solicitando parada da extracao/monitoramento...")
        self._encerrar_driver_ativo()

    def fechar_aplicacao(self):
        self._salvar_grupos_monitorados()
        self.cancelar_tudo(atualizar_status=False)
        try:
            if self.dashboard_window and self.dashboard_window.winfo_exists():
                self.dashboard_window.destroy()
        except Exception:
            pass
        self.root.destroy()

    def iniciar(self):
        if self.rodando:
            return
        faixas = self.get_faixas_selecionadas()
        if not faixas:
            messagebox.showwarning("Atencao", "Selecione ao menos uma faixa.")
            return
        self.cancelar_execucao.clear()
        self.rodando = True
        self.btn.configure(state="disabled", text="Extraindo...")
        self.btn_tudo.configure(state="disabled")
        self.btn_iniciar_monitor.configure(state="disabled")
        self.btn_monitor_dashboard.configure(state="disabled")
        self.btn_excel.configure(state="disabled")
        self.log_box.configure(state="normal")
        self.log_box.delete("1.0", "end")
        self.log_box.configure(state="disabled")
        self.set_status("Extraindo...")
        if faixas == ["__TODOS__"]:
            threading.Thread(target=self.executar_tudo, daemon=True).start()
        else:
            threading.Thread(
                target=self.executar,
                args=(self.tipo_bem.get(), faixas, self.get_grupos_monitorados()),
                daemon=True
            ).start()

    def iniciar_monitoramento(self):
        if self.monitorando:
            return
        faixas = self.get_faixas_selecionadas()
        grupos = self.get_grupos_monitorados()
        if not faixas:
            messagebox.showwarning("Atencao", "Selecione ao menos uma faixa.")
            return
        if not grupos:
            messagebox.showwarning("Atencao", "Adicione ao menos um grupo.")
            return
        self.cancelar_execucao.clear()
        self.monitorando = True
        self.parar_monitor.clear()
        self.btn_iniciar_monitor.configure(state="disabled")
        self.btn_parar_monitor.configure(state="normal")
        self.btn_monitor_dashboard.configure(state="disabled")
        self.log_box.configure(state="normal")
        self.log_box.delete("1.0", "end")
        self.log_box.configure(state="disabled")
        self.set_status("Monitorando...")
        modo_tudo = faixas == ["__TODOS__"]
        threading.Thread(
            target=self.loop_monitoramento,
            args=(self.tipo_bem.get(), faixas, grupos, modo_tudo),
            daemon=True
        ).start()

    def parar_monitoramento(self):
        self.cancelar_tudo()
        self.btn_iniciar_monitor.configure(state="normal")
        self.btn_parar_monitor.configure(state="disabled")
        self.btn_monitor_dashboard.configure(state="normal", text="Monitorar + Dashboard")
        self.label_proximo.configure(text="Monitoramento pausado.")
        self.set_status("Pronto")
        self.log("Monitoramento parado.")

    def loop_monitoramento(self, tipo, faixas, grupos, modo_tudo=False):
        rodada = 1
        while not self.parar_monitor.is_set() and not self.cancelar_execucao.is_set():
            self.log(f"Verificacao #{rodada} - {datetime.now().strftime('%H:%M:%S')}")
            if modo_tudo:
                grupos_encontrados = self._executar_extracao_tudo(grupos)
            else:
                grupos_encontrados = self._executar_extracao(tipo, faixas, grupos)
            if grupos_encontrados and self.comportamento_var.get() == "parar":
                self.root.after(0, self.parar_monitoramento)
                break
            if self.parar_monitor.is_set() or self.cancelar_execucao.is_set():
                break
            intervalo = self.get_intervalo_segundos()
            for restante in range(intervalo, 0, -1):
                if self.parar_monitor.is_set() or self.cancelar_execucao.is_set():
                    break
                mins, secs = divmod(restante, 60)
                self.root.after(0, lambda m=mins, s=secs: self.label_proximo.configure(
                    text=f"Proxima verificacao em {m:02d}:{s:02d}"
                ))
                time.sleep(1)
            rodada += 1
        self.monitorando = False
        self.root.after(0, lambda: self.btn_monitor_dashboard.configure(state="normal", text="Monitorar + Dashboard"))

    def iniciar_monitoramento_com_dashboard(self):
        if self.monitorando:
            self.abrir_dashboard()
            return
        dashboard = self.abrir_dashboard()
        if dashboard is None:
            return
        self.btn_monitor_dashboard.configure(text="Dashboard aberto...")
        self.iniciar_monitoramento()

    def abrir_dashboard(self):
        if self.dashboard_window and self.dashboard_window.winfo_exists():
            self.dashboard_window.deiconify()
            self.dashboard_window.lift()
            return self.dashboard_window

        try:
            try:
                from dashboardtela import DashboardTV
            except Exception:
                dashboard_path = None
                for nome in ("dashboardtela.py", "dashboardtela (4).py"):
                    caminho = os.path.join(PASTA_BASE, nome)
                    if os.path.exists(caminho):
                        dashboard_path = caminho
                        break
                if not dashboard_path:
                    raise FileNotFoundError("dashboardtela.py nao encontrado")
                spec = importlib.util.spec_from_file_location("dashboardtela_mod", dashboard_path)
                modulo = importlib.util.module_from_spec(spec)
                spec.loader.exec_module(modulo)
                DashboardTV = modulo.DashboardTV

            janela_master = self.root.master if isinstance(self.root, tk.Toplevel) else self.root
            self.dashboard_window = DashboardTV(master=janela_master, keep_on_top=True, exit_on_close=False)
            return self.dashboard_window
        except Exception as err:
            self.btn_monitor_dashboard.configure(state="normal", text="Monitorar + Dashboard")
            messagebox.showerror("Dashboard", f"Nao foi possivel abrir o dashboard:\n{err}")
            return None

    def _executar_extracao(self, tipo, faixas, grupos_monitorados):
        MAX_TENTATIVAS = 3
        grupos_encontrados = []
        for tentativa in range(1, MAX_TENTATIVAS + 1):
            if self.cancelar_execucao.is_set():
                self.log("Extracao cancelada antes de iniciar a tentativa.")
                break
            driver = None
            try:
                self.log(f"Abrindo navegador... (tentativa {tentativa}/{MAX_TENTATIVAS})")
                driver = configurar_navegador(exibir_janela=self.mostrar_navegador_var.get())
                self._registrar_driver(driver)
                wait = WebDriverWait(driver, TIMEOUT)
                fazer_login(driver, wait, self.log, self.config)
                abas = []
                for faixa in faixas:
                    if self.cancelar_execucao.is_set() or self.parar_monitor.is_set():
                        self.log("Cancelamento solicitado. Encerrando extracao atual...")
                        break
                    cabecalho, linhas = acessar_simulador_faixa(driver, wait, self.log, tipo, faixa)
                    if grupos_monitorados:
                        for linha in linhas:
                            if linha and linha[0].strip().upper() in grupos_monitorados:
                                grupos_encontrados.append({"grupo": linha[0].strip(), "faixa": faixa})
                    nome_aba = faixa.replace("De R$ ", "").replace(" ate R$ ", "-").replace(" até R$ ", "-").replace(".", "").replace(",00", "")
                    abas.append((nome_aba, cabecalho, linhas))
                if self.cancelar_execucao.is_set() or self.parar_monitor.is_set():
                    break
                grupos_encontrados = deduplicar_grupos_encontrados(grupos_encontrados)
                if abas:
                    salvar_excel(abas, self.log)
                    total = sum(len(d) for _, _, d in abas)
                    self.log(f"Concluido! {total} grupos extraidos.")
                    self.root.after(0, lambda: self.btn_excel.configure(state="normal"))
                if grupos_encontrados:
                    gl = "\n".join(f"  - {g['grupo']}  ({g['faixa']})" for g in grupos_encontrados)
                    self.log(f"ALERTA: {', '.join(g['grupo'] for g in grupos_encontrados)}")
                    self.root.after(0, lambda g=gl: messagebox.showwarning(
                        "Grupos encontrados!", f"Os seguintes grupos estao disponiveis:\n\n{g}"
                    ))
                    resultado = enviar_email(grupos_encontrados, self.config)
                    self.log("E-mail enviado!" if resultado is True else f"Falha no e-mail: {resultado}")
                elif grupos_monitorados:
                    self.log("Nenhum grupo monitorado encontrado.")
                break
            except FileNotFoundError as err:
                msg = str(err)
                self.log(f"Erro: {msg}")
                self.root.after(0, lambda m=msg: messagebox.showerror("Driver nao encontrado", m))
                break
            except Exception as err:
                if self._erro_fatal_sem_retry(err):
                    self.log("Sessao do navegador encerrada. Retentativa automatica cancelada.")
                    break
                self.log(f"Erro (tentativa {tentativa}/{MAX_TENTATIVAS}): {str(err)[:120]}")
                if tentativa == MAX_TENTATIVAS:
                    self.log("Numero maximo de tentativas atingido.")
            finally:
                if driver:
                    try: driver.quit()
                    except: pass
                if self.driver_ativo is driver:
                    self.driver_ativo = None
            if tentativa < MAX_TENTATIVAS:
                if self.cancelar_execucao.is_set():
                    break
                self.log("Aguardando 10s para nova tentativa...")
                time.sleep(10)
        return grupos_encontrados

    def executar(self, tipo, faixas, grupos):
        self._executar_extracao(tipo, faixas, grupos)
        if self.cancelar_execucao.is_set():
            self.log("Extracao interrompida pelo usuario.")
        elif not grupos:
            messagebox.showinfo("Concluido", "Extracao finalizada!")
        self.rodando = False
        self.set_status("Pronto")
        self.btn.configure(state="normal", text="EXTRAIR UMA VEZ")
        self.btn_tudo.configure(state="normal")
        self.btn_iniciar_monitor.configure(state="normal")
        self.btn_monitor_dashboard.configure(state="normal", text="Monitorar + Dashboard")

    def iniciar_tudo(self):
        if self.rodando:
            return
        self.cancelar_execucao.clear()
        self.rodando = True
        self.btn.configure(state="disabled")
        self.btn_tudo.configure(state="disabled", text="Extraindo tudo...")
        self.btn_iniciar_monitor.configure(state="disabled")
        self.btn_monitor_dashboard.configure(state="disabled")
        self.btn_excel.configure(state="disabled")
        self.log_box.configure(state="normal")
        self.log_box.delete("1.0", "end")
        self.log_box.configure(state="disabled")
        self.set_status("Extraindo tudo...")
        threading.Thread(target=self.executar_tudo, daemon=True).start()

    def _executar_extracao_tudo(self, grupos_monitorados=None):
        MAX_TENTATIVAS = 3
        grupos_encontrados = []
        tipos_base = ["Automoveis", "Veiculos pesados", "Imoveis"]
        for tentativa in range(1, MAX_TENTATIVAS + 1):
            if self.cancelar_execucao.is_set():
                self.log("Extracao completa cancelada antes de iniciar a tentativa.")
                break
            driver = None
            try:
                self.log(f"Abrindo navegador... (tentativa {tentativa}/{MAX_TENTATIVAS})")
                driver = configurar_navegador(exibir_janela=self.mostrar_navegador_var.get())
                self._registrar_driver(driver)
                wait = WebDriverWait(driver, TIMEOUT)
                fazer_login(driver, wait, self.log, self.config)
                for tipo_chave in tipos_base:
                    if self.cancelar_execucao.is_set() or self.parar_monitor.is_set():
                        self.log("Cancelamento solicitado. Encerrando extracao atual...")
                        break
                    faixas = FAIXAS_REAIS[tipo_chave]
                    tipo_real = TIPOS_REAIS[tipo_chave]
                    self.log(f"=== {tipo_real} ({len(faixas)} faixas) ===")
                    abas = []
                    for faixa in faixas:
                        if self.cancelar_execucao.is_set() or self.parar_monitor.is_set():
                            break
                        cabecalho, linhas = acessar_simulador_faixa(driver, wait, self.log, tipo_chave, faixa)
                        if grupos_monitorados:
                            for linha in linhas:
                                if linha and linha[0].strip().upper() in grupos_monitorados:
                                    grupos_encontrados.append({"grupo": linha[0].strip(), "faixa": faixa, "tipo": tipo_real})
                        nome_aba = faixa.replace("De R$ ", "").replace(" até R$ ", "-").replace(".", "").replace(",00", "")
                        abas.append((nome_aba, cabecalho, linhas))
                    if self.cancelar_execucao.is_set() or self.parar_monitor.is_set():
                        break
                    nome_arquivo = os.path.join(PASTA_BASE, f"vagas_{tipo_chave.lower().replace(' ', '_')}.xlsx")
                    salvar_excel(abas, self.log, arquivo=nome_arquivo)
                    total = sum(len(d) for _, _, d in abas)
                    self.log(f"=== {tipo_real}: {total} grupos ===")
                grupos_encontrados = deduplicar_grupos_encontrados(grupos_encontrados)
                if self.cancelar_execucao.is_set() or self.parar_monitor.is_set():
                    break
                self.log("Extracao completa finalizada!")
                self.root.after(0, lambda: self.btn_excel.configure(state="normal"))
                if grupos_encontrados:
                    gl = "\n".join(f"  - {g['grupo']} ({g['tipo']} | {g['faixa']})" for g in grupos_encontrados)
                    self.log(f"ALERTA: {', '.join(g['grupo'] for g in grupos_encontrados)}")
                    self.root.after(0, lambda g=gl: messagebox.showwarning(
                        "Grupos encontrados!", f"Os seguintes grupos estao disponiveis:\n\n{g}"
                    ))
                    resultado = enviar_email(grupos_encontrados, self.config)
                    self.log("E-mail enviado!" if resultado is True else f"Falha no e-mail: {resultado}")
                elif grupos_monitorados:
                    self.log("Nenhum grupo monitorado encontrado.")
                break
            except Exception as err:
                if self._erro_fatal_sem_retry(err):
                    self.log("Sessao do navegador encerrada. Retentativa automatica cancelada.")
                    break
                self.log(f"Erro (tentativa {tentativa}/{MAX_TENTATIVAS}): {str(err)[:120]}")
                if tentativa == MAX_TENTATIVAS:
                    self.log("Numero maximo de tentativas atingido.")
            finally:
                if driver:
                    try: driver.quit()
                    except: pass
                if self.driver_ativo is driver:
                    self.driver_ativo = None
            if tentativa < MAX_TENTATIVAS:
                if self.cancelar_execucao.is_set():
                    break
                self.log("Aguardando 10s para nova tentativa...")
                time.sleep(10)
        return grupos_encontrados

    def executar_tudo(self):
        grupos = self.get_grupos_monitorados()
        self._executar_extracao_tudo(grupos_monitorados=grupos)
        if self.cancelar_execucao.is_set():
            self.log("Extracao completa interrompida pelo usuario.")
        self.rodando = False
        self.set_status("Pronto")
        self.root.after(0, lambda: self.btn.configure(state="normal"))
        self.root.after(0, lambda: self.btn_tudo.configure(state="normal", text="EXTRAIR TUDO  (Auto + Pesados + Imoveis)"))
        self.root.after(0, lambda: self.btn_iniciar_monitor.configure(state="normal"))
        self.root.after(0, lambda: self.btn_monitor_dashboard.configure(state="normal", text="Monitorar + Dashboard"))

    def abrir_pasta(self):
        os.startfile(PASTA_BASE)


# ============================================================
# MAIN
# ============================================================

if __name__ == "__main__":
    modo_auto = "--auto" in sys.argv
    root = tk.Tk()
    app = App(root)
    if modo_auto:
        def iniciar_auto():
            faixas = app.get_faixas_selecionadas()
            grupos = app.get_grupos_monitorados()
            if not faixas:
                app.log("AVISO: Nenhuma faixa selecionada.")
                return
            if not grupos:
                app.log("AVISO: Nenhum grupo para monitorar.")
                return
            app.log("Modo automatico ativado...")
            app.iniciar_monitoramento()
        root.after(3000, iniciar_auto)
    root.mainloop()
